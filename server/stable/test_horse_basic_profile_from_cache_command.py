"""A036 独立合成测试：只在 ROOT 精确分配的官方 Linux 执行器内运行。"""
import csv
import hashlib
import io
import json
import os
import stat
import subprocess
import sys
import tempfile
from contextlib import ExitStack
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import SimpleTestCase
from openpyxl import load_workbook

from stable.services.horse_cache_reuse import plan_cache_reuse
from stable.services.horse_source_cache_reuse_adapter import adapt_hkjc_source_cache
from stable.services.p0_horse_completion_adapters import (
    P0HorseCompletionRequest, REGION_ADAPTERS,
)
from stable.services.p0_horse_completion_review import build_batch_review_workbook

COMMAND = "horse_basic_profile_from_cache"
FIXTURE = Path(__file__).parent / "fixtures/h03_prepare_command/hkjc_synthetic.json"
FIELDS = ("country", "sex", "color", "birth_date", "owner_name", "trainer_name", "breeder_name")
DECISIONS = ("basic_profile_decision", "pedigree_decision", "race_records_decision",
             "major_wins_decision", "reviewer_id", "reviewed_at")
SERVER = Path(__file__).resolve().parents[1]


def encoded(value):
    return (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode()


def sha(data):
    return hashlib.sha256(data).hexdigest()


def snapshot():
    return {
        "schema_version": 1, "as_of": "2026-10-03",
        "snapshot_at": "2026-10-03T00:00:00+00:00", "input_complete": True,
        "policy_sha": "a" * 64, "jg1_scope": "unresolved",
        "events": [{"key": "e1", "region": "hong_kong", "local_date": "2026-10-03",
                    "date_precision": "exact", "grade": "G1", "grade_verified": True,
                    "hk_local_g1": False, "current_racecard": None,
                    "current_result": None, "known_roster_count": 1}],
        "participations": [{"ref": "r1", "event_key": "e1", "participant_key": "p1",
                            "horse_key": "hkjc:HK-001", "revision_ref": None,
                            "kind": "legacy", "status": "finished", "start_evidence": True,
                            "birth_year": 2020}],
        "identities": [{"horse_key": "hkjc:HK-001", "profile_id": 1,
                        "status": "verified", "evidence_sha": "b" * 64,
                        "verified_at": "2026-10-02T00:00:00+00:00"}],
        "layers": [{"target_key": "profile:1", "cache": "present", "staging": "matched",
                    "profile_exists": True, "public_state": "unpublished",
                    "starts": 0, "incomplete_modules": 1}],
        "news": [], "seasons": [],
    }


# Installed before django.setup in a real CLI subprocess, not merely around handle().
CLI_BOOTSTRAP = r'''
import json, os, runpy, socket, sys
from pathlib import Path
from unittest.mock import patch
from contextlib import ExitStack
import dotenv
dotenv.load_dotenv = lambda *args, **kwargs: False
sys.path.insert(0, sys.argv[1])
mode, args = sys.argv[2], sys.argv[3:]
trips = []
def forbidden(*args, **kwargs):
    trips.append("database_or_network")
    raise AssertionError("A036 all-alias DB/network tripwire")
code, ready = 0, False
with ExitStack() as stack:
    for method in ("connect", "ensure_connection", "cursor", "_cursor"):
        stack.enter_context(patch("django.db.backends.base.base.BaseDatabaseWrapper." + method,
                                  side_effect=forbidden))
    for method in ("connect", "connect_ex"):
        stack.enter_context(patch("socket.socket." + method, side_effect=forbidden))
    from django.conf import settings
    settings.DATABASES = {name: {
        "ENGINE": "django.db.backends.postgresql", "NAME": "synthetic_no_access",
        "USER": "synthetic_no_access", "PASSWORD": "synthetic-only",
        "HOST": "203.0.113.10" if mode == "production_style" else "127.0.0.1",
        "PORT": "5432", "CONN_MAX_AGE": 0,
    } for name in ("default", "replica")}
    os.environ["DATABASE_URL"] = "postgres://synthetic:synthetic@203.0.113.11/never_connect"
    try:
        import django
        django.setup()
        ready = True
        sys.argv = [str(Path(sys.path[0]) / "manage.py"), *args]
        runpy.run_path(sys.argv[0], run_name="__main__")
    except SystemExit as exc:
        code = exc.code if isinstance(exc.code, int) else 1
    except BaseException:
        code = 7
    finally:
        sys.stderr.write("A036_GUARD=" + json.dumps({"setup_completed": ready,
            "aliases": sorted(settings.DATABASES), "trips": trips}) + "\n")
sys.exit(code)
'''


class PrepareCommandTests(SimpleTestCase):
    databases = set()

    def setUp(self):
        super().setUp()
        self.tmp = tempfile.TemporaryDirectory(prefix="a036-synthetic-")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.input_root = self.root / "input"
        self.output_root = self.root / "output"
        self.input_root.mkdir(mode=0o700)
        self.output_root.mkdir(mode=0o700)
        self.raw = FIXTURE.read_bytes()
        self.packet = {
            "schema_version": "h03-prepare-input.v1", "source_file": "cache.json",
            "cache_ref": "cache:a036", "source_ref": "fixture:a036-hkjc",
            "snapshot": snapshot(), "entity_versions": {"profile:1": "a036-v1"},
            "as_of": "2026-10-03T01:00:00+00:00", "max_age_seconds": 77 * 86400 + 3600,
            "profile_baseline": "2026-10-02T00:00:00+00:00",
        }
        self.bind_inputs()
        self.guards = ExitStack()
        self.addCleanup(self.guards.close)
        for target in (
            "django.db.backends.base.base.BaseDatabaseWrapper.connect",
            "django.db.backends.base.base.BaseDatabaseWrapper.ensure_connection",
            "django.db.backends.base.base.BaseDatabaseWrapper.cursor",
            "django.db.backends.base.base.BaseDatabaseWrapper._cursor",
            "socket.socket.connect", "socket.socket.connect_ex", "urllib.request.urlopen",
            "stable.services.horse_basic_profile_from_cache.apply_basic_profile_from_cache",
            "stable.services.horse_profiles.upsert_race_record",
            "stable.services.horse_profile_publish.auto_publish_profiles",
            "stable.services.p0_horse_completion_adapters.run_p0_horse_completion_adapter",
            "stable.services.p0_horse_completion_source_clients._HKJCClient._fetch",
        ):
            self.guards.enter_context(patch(target, side_effect=AssertionError("A036 DB/network/write forbidden")))

    def bind_inputs(self, packet_bytes=None):
        (self.input_root / "cache.json").write_bytes(self.raw)
        packet_bytes = encoded(self.packet) if packet_bytes is None else packet_bytes
        (self.input_root / "packet.json").write_bytes(packet_bytes)
        # Synthetic caller independently binds both streams; not real source approval.
        self.kw = {"input_root": str(self.input_root), "input": "packet.json",
                   "expected_input_sha256": sha(packet_bytes), "expected_source_sha256": sha(self.raw),
                   "output_root": str(self.output_root), "output_dir": "pending"}

    def preflight(self):
        source = json.loads(self.raw)
        record = adapt_hkjc_source_cache(self.raw, expected_sha256=sha(self.raw),
                                         ref=self.packet["cache_ref"], source_ref=self.packet["source_ref"])
        plan = plan_cache_reuse(self.packet["snapshot"], [record], self.packet["entity_versions"],
                                as_of=self.packet["as_of"], max_age_seconds=self.packet["max_age_seconds"])
        self.assertEqual([d["status"] for d in plan["decisions"]], ["reusable"])
        self.assertEqual(len(plan["candidates"]), 1)
        request = P0HorseCompletionRequest(
            candidate_key=plan["candidates"][0]["idempotency_key"], region="hong_kong",
            horse_name=source["identity"]["horse_name"], source_url=source["source"]["url"],
            external_horse_id=source["source"]["external_horse_id"], candidate_source_name="hkjc")
        canonical = REGION_ADAPTERS["hong_kong"].normalize(source, request)
        self.assertEqual(canonical["basic_profile"], source["basic_profile"])
        self.assertEqual(canonical["raw_payload"], source)
        self.assertEqual(canonical["failure_reason"], [])
        # Prove real consumer and fixture are valid before asserting absent command behavior.
        check = self.root / "preflight"
        check.mkdir(exist_ok=True)
        (check / "combined_candidates.jsonl").write_text(json.dumps(canonical, ensure_ascii=False) + "\n")
        path = build_batch_review_workbook(manifest={"status": "local_pending_review", "horses": []},
                                          artifact_dir=check, output_path=check / "view.xlsx")
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            self.assertIn("中国香港", wb.sheetnames)
            self.assertEqual(wb["中国香港"].cell(2, 2).value, canonical["candidate_key"])
        finally:
            wb.close()
        return canonical, plan

    def invoke(self, **changes):
        output = io.StringIO()
        call_command(COMMAND, stdout=output, stderr=io.StringIO(), **{**self.kw, **changes})
        text = output.getvalue()
        result = json.loads(text)
        return result, text

    def assert_prepared(self):
        result, text = self.invoke()
        self.assertEqual(result["status"], "prepared", "normal command stub has no prepare behavior")
        self.assertEqual(result["reason"], "local_pending_review")
        self.assertIs(result["reviewed"], False)
        self.assertNotIn("applied", result)
        return result, text

    def cli(self, extra=(), mode="ordinary", directory="cli-pending"):
        options = {**self.kw, "output_dir": directory}
        args = [COMMAND]
        for key, value in options.items():
            args.extend(["--" + key.replace("_", "-"), value])
        env = {k: os.environ[k] for k in ("PATH", "HOME", "TMPDIR") if k in os.environ}
        env.update(DJANGO_SETTINGS_MODULE="app.settings", PYTHONDONTWRITEBYTECODE="1",
                   DB_ENGINE="postgres", POSTGRES_DB="synthetic_no_access", POSTGRES_USER="synthetic_no_access",
                   POSTGRES_PASSWORD="synthetic-only", POSTGRES_HOST="203.0.113.10",
                   CELERY_BROKER_URL="memory://", CELERY_RESULT_BACKEND="cache+memory://")
        run = subprocess.run([sys.executable, "-c", CLI_BOOTSTRAP, str(SERVER), mode, *args, *extra],
                             env=env, cwd=SERVER, capture_output=True, text=True, timeout=30)
        marker = [line for line in run.stderr.splitlines() if line.startswith("A036_GUARD=")]
        self.assertEqual(len(marker), 1, "CLI startup did not reach guard receipt")
        guard = json.loads(marker[0].split("=", 1)[1])
        self.assertTrue(guard["setup_completed"], "setup/import errors are not business RED")
        self.assertEqual(guard["aliases"], ["default", "replica"])
        self.assertEqual(guard["trips"], [], "implicit DB/network access")
        return run

    def blocked(self, reason, **changes):
        with self.assertRaises(CommandError) as caught:
            self.invoke(**changes)
        self.assertEqual(str(caught.exception), reason)
        self.assertFalse((self.output_root / "pending").exists())

    def test_prepare_real_pipeline_pending_outputs_and_visible_seven_fields(self):
        canonical, plan = self.preflight()
        with patch("stable.services.horse_source_cache_reuse_adapter.adapt_hkjc_source_cache", wraps=adapt_hkjc_source_cache) as adapter, \
             patch("stable.services.horse_cache_reuse.plan_cache_reuse", wraps=plan_cache_reuse) as planner, \
             patch("stable.services.p0_horse_completion_review.build_batch_review_workbook", wraps=build_batch_review_workbook) as workbook:
            result, _ = self.assert_prepared()  # Sole initial requested RED: not_implemented != prepared.
            self.assertEqual(adapter.call_count, 1)
            self.assertEqual(planner.call_count, 1)
            self.assertEqual(workbook.call_count, 1)
        self.assertEqual(result["basic_profile"], {k: json.loads(self.raw)["basic_profile"][k] for k in FIELDS})
        out = self.output_root / "pending"
        self.assertEqual(Path(result["artifact_dir"]), out)
        candidates = [json.loads(line) for line in (out / "combined_candidates.jsonl").read_text().splitlines()]
        self.assertEqual(len(candidates), 1)
        candidate = candidates[0]
        self.assertIs(candidate["reviewed"], False)
        for field, value in canonical.items():
            self.assertEqual(candidate[field], value)
        manifest = json.loads((out / "manifest.json").read_text())
        self.assertEqual(manifest["status"], "local_pending_review")
        self.assertEqual(manifest["input_sha256"], self.kw["expected_input_sha256"])
        self.assertEqual(manifest["source_sha256"], self.kw["expected_source_sha256"])
        self.assertEqual(manifest["h02_sha256"], plan["content_sha256"])
        self.assertEqual(manifest["candidate_idempotency_key"], plan["candidates"][0]["idempotency_key"])
        self.assertNotIn("reviewed_input_sha256", manifest)
        with (out / "review.csv").open(newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["reviewed"].lower(), "false")
        self.assertTrue(all(rows[0][key] == "" for key in DECISIONS))
        wb = load_workbook(out / "review.xlsx", read_only=True)
        try:
            self.assertEqual(wb["中国香港"].cell(2, 2).value, candidate["candidate_key"])
            self.assertEqual(wb["汇总"].cell(3, 2).value, "local_pending_review")
        finally:
            wb.close()

    def test_cli_startup_and_workbook_have_zero_database_access(self):
        self.preflight()
        for mode in ("ordinary", "production_style"):
            with self.subTest(mode=mode):
                run = self.cli(mode=mode, directory=mode)
                self.assertEqual(run.returncode, 0, "command/import error is not valid preparation")
                result = json.loads(run.stdout)
                self.assertEqual(result["status"], "prepared")
                wb = load_workbook(self.output_root / mode / "review.xlsx", read_only=True)
                try:
                    self.assertIn("中国香港", wb.sheetnames)
                    self.assertEqual(wb["中国香港"].cell(2, 3).value, "HARBOUR TEST")
                finally:
                    wb.close()

    def test_parser_rejects_apply_and_commit_without_database_access(self):
        self.preflight()
        for flag in ("--apply", "--commit"):
            with self.subTest(flag=flag):
                run = self.cli(extra=[flag], directory="forbidden-write")
                self.assertEqual(run.returncode, 2)
                self.assertIn("unrecognized arguments", run.stderr)
                self.assertIn(flag, run.stderr)
                self.assertFalse((self.output_root / "forbidden-write").exists())

    def test_trust_hash_identity_and_freshness_fail_closed(self):
        self.preflight()
        for key, reason in (("expected_input_sha256", "input_hash_mismatch"),
                            ("expected_source_sha256", "content_hash_mismatch")):
            with self.subTest(key=key):
                self.blocked(reason, **{key: "f" * 64})
        for key, reason in (("expected_input_sha256", "input_hash_unbound"),
                            ("expected_source_sha256", "content_hash_unbound")):
            with self.subTest(missing=key): self.blocked(reason, **{key: ""})
        cached = self.input_root / "cache.json"
        cached.write_bytes(self.raw + b" ")
        self.blocked("content_hash_mismatch")
        self.bind_inputs()
        original = deepcopy(self.packet)
        cases = ("retired", "conflict", "stale", "approval", "partial", "source_identity")
        for case in cases:
            with self.subTest(case=case):
                self.packet = deepcopy(original); self.raw = FIXTURE.read_bytes()
                reason = "cache_not_reusable"
                if case == "retired": self.packet["snapshot"]["identities"][0]["status"] = "retired"
                if case == "conflict":
                    proof = deepcopy(self.packet["snapshot"]["identities"][0]); proof["profile_id"] = 2
                    self.packet["snapshot"]["identities"].append(proof)
                if case == "stale": self.packet["max_age_seconds"] = 3600
                if case == "approval": self.packet["reviewed"] = True; reason = "input_schema"
                if case in ("partial", "source_identity"):
                    raw = json.loads(self.raw)
                    if case == "partial": raw["basic_profile"].pop("trainer_name"); reason = "cache_validation"
                    else: raw["source"]["external_horse_id"] = "other"; reason = "provider_identity_url"
                    self.raw = encoded(raw)
                self.bind_inputs()
                self.blocked(reason)
                self.assertEqual((self.input_root / "cache.json").read_bytes(), self.raw)

    def test_path_and_json_limits_fail_closed_without_overwrite(self):
        self.preflight()
        self.blocked("input_path", input="../packet.json")
        self.blocked("output_path", output_dir="../outside")
        (self.input_root / "packet-link.json").symlink_to(self.input_root / "packet.json")
        self.blocked("input_path", input="packet-link.json")
        for name, target in (("input-link", self.input_root), ("output-link", self.output_root)):
            link = self.root / name; link.symlink_to(target, target_is_directory=True)
            self.blocked("input_path" if name == "input-link" else "output_path",
                         **{"input_root" if name == "input-link" else "output_root": str(link)})
        original = deepcopy(self.packet)
        self.packet["source_file"] = "../cache.json"; self.bind_inputs(); self.blocked("input_path")
        cache_link = self.input_root / "cache-link.json"; cache_link.symlink_to(self.input_root / "cache.json")
        self.packet["source_file"] = cache_link.name; self.bind_inputs(); self.blocked("input_path")
        parent = self.input_root / "parent-link"
        parent.symlink_to(self.input_root, target_is_directory=True)
        self.packet["source_file"] = "parent-link/cache.json"; self.bind_inputs(); self.blocked("input_path")
        self.packet = deepcopy(original)
        folder = self.input_root / "folder"; folder.mkdir()
        self.packet["source_file"] = "folder"; self.bind_inputs(); self.blocked("input_path")
        self.packet = deepcopy(original)
        for raw in (b'{"x":1,"x":2}', b'{"x":NaN}', b'{"x":Infinity}', b'[' * 33 + b'0' + b']' * 33):
            with self.subTest(raw=raw[:24]):
                self.bind_inputs(packet_bytes=raw); self.blocked("input_json")
        self.bind_inputs(packet_bytes=b' ' * 131073); self.blocked("input_size")
        self.raw = b' ' * 131073; self.bind_inputs(); self.blocked("cache_size")
        self.raw = FIXTURE.read_bytes(); self.bind_inputs()
        out = self.output_root / "pending"; out.mkdir(); sentinel = out / "other-owner.txt"; sentinel.write_bytes(b"preserve")
        with self.assertRaises(CommandError) as caught: self.invoke()
        self.assertEqual(str(caught.exception), "output_exists")
        self.assertEqual(sentinel.read_bytes(), b"preserve")
        self.assertEqual(list(out.iterdir()), [sentinel])

    def test_display_canonical_separation_and_downstream_failure_cleanup(self):
        raw = json.loads(self.raw)
        raw["identity"]["horse_name"] = "=SYNTHETIC"
        raw["aliases"][0]["name"] = "=SYNTHETIC"
        raw["basic_profile"]["owner_name"] = "=1+1\t\x1b[31m"
        self.raw = encoded(raw); self.bind_inputs()
        canonical, _ = self.preflight()
        _, stdout = self.assert_prepared()
        self.assertNotIn("\x1b", stdout); self.assertNotIn("\t", stdout)
        out = self.output_root / "pending"
        candidate = json.loads((out / "combined_candidates.jsonl").read_text())
        self.assertEqual(candidate["horse_name"], "=SYNTHETIC")
        self.assertEqual(candidate["raw_payload"], raw)
        self.assertEqual(candidate["basic_profile"], canonical["basic_profile"])
        with (out / "review.csv").open(newline="") as stream:
            self.assertEqual(list(csv.DictReader(stream))[0]["horse_name"], "'=SYNTHETIC")
        wb = load_workbook(out / "review.xlsx", read_only=True, data_only=False)
        try:
            cell = wb["中国香港"].cell(2, 3)
            self.assertEqual(cell.value, "'=SYNTHETIC"); self.assertNotEqual(cell.data_type, "f")
        finally:
            wb.close()
        self.assertEqual(stat.S_IMODE(out.stat().st_mode), 0o700)
        self.assertTrue(all(stat.S_IMODE(p.stat().st_mode) == 0o600 for p in out.iterdir()))
        with patch("stable.services.p0_horse_completion_review.build_batch_review_workbook", side_effect=OSError("synthetic IO fault")):
            with self.assertRaises(CommandError) as caught: self.invoke(output_dir="io-failure")
        self.assertEqual(str(caught.exception), "artifact_build_failed")
        self.assertFalse((self.output_root / "io-failure").exists())
        self.assertTrue((out / "review.xlsx").exists())
        self.assertEqual(sorted(p.name for p in self.output_root.iterdir()), ["pending"])
